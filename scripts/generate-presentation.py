#!/usr/bin/env python3
"""
Generate a PowerPoint presentation for the UK Retail Market Financial Newsletter.

Usage:
    python scripts/generate-presentation.py
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Brand colours
NAVY = RGBColor(0x0B, 0x1D, 0x3A)
DARK_BLUE = RGBColor(0x14, 0x2D, 0x5E)
ACCENT_BLUE = RGBColor(0x1B, 0x6B, 0xB0)
LIGHT_BLUE = RGBColor(0xD6, 0xEA, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GREY = RGBColor(0xF2, 0xF2, 0xF2)
MID_GREY = RGBColor(0x7F, 0x8C, 0x8D)
DARK_GREY = RGBColor(0x2C, 0x3E, 0x50)
GREEN = RGBColor(0x27, 0xAE, 0x60)
RED = RGBColor(0xC0, 0x39, 0x2B)
AMBER = RGBColor(0xF3, 0x9C, 0x12)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def add_background(slide, colour=NAVY):
    """Fill the slide background with a solid colour."""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = colour


def add_shape_bg(slide, left, top, width, height, colour):
    """Add a filled rectangle (no border) as a background panel."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = colour
    shape.line.fill.background()
    return shape


def add_textbox(slide, left, top, width, height, text, font_size=18,
                colour=WHITE, bold=False, alignment=PP_ALIGN.LEFT, font_name="Calibri"):
    """Add a simple text box."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = colour
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_table(slide, left, top, width, row_height, data, col_widths=None):
    """Add a table from a 2D list. First row is treated as header."""
    rows, cols = len(data), len(data[0])
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, Inches(row_height * rows))
    table = table_shape.table

    if col_widths:
        for i, w in enumerate(col_widths):
            table.columns[i].width = Inches(w)

    for r, row_data in enumerate(data):
        for c, cell_text in enumerate(row_data):
            cell = table.cell(r, c)
            cell.text = str(cell_text)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE

            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(11)
                paragraph.font.name = "Calibri"
                if r == 0:
                    paragraph.font.bold = True
                    paragraph.font.color.rgb = WHITE
                    paragraph.alignment = PP_ALIGN.CENTER
                else:
                    paragraph.font.color.rgb = DARK_GREY
                    paragraph.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER

            if r == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = DARK_BLUE
            elif r % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = LIGHT_GREY
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE

    return table_shape


def section_header(slide, text):
    """Add the blue accent bar + section title at the top of a content slide."""
    add_shape_bg(slide, Inches(0), Inches(0), SLIDE_W, Inches(1.1), NAVY)
    add_shape_bg(slide, Inches(0), Inches(1.1), SLIDE_W, Inches(0.06), ACCENT_BLUE)
    add_textbox(slide, Inches(0.6), Inches(0.25), Inches(12), Inches(0.7),
                text, font_size=28, colour=WHITE, bold=True)


# ── SLIDE 1: Title ──────────────────────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
add_background(slide, NAVY)
add_shape_bg(slide, Inches(0), Inches(2.8), SLIDE_W, Inches(2.6), DARK_BLUE)
add_textbox(slide, Inches(0.8), Inches(3.0), Inches(11.5), Inches(1.0),
            "UK Retail Market", font_size=44, colour=WHITE, bold=True,
            alignment=PP_ALIGN.CENTER)
add_textbox(slide, Inches(0.8), Inches(3.8), Inches(11.5), Inches(0.7),
            "Monthly Financial Newsletter", font_size=32, colour=LIGHT_BLUE,
            alignment=PP_ALIGN.CENTER)
add_textbox(slide, Inches(0.8), Inches(4.6), Inches(11.5), Inches(0.5),
            "March 2026 Issue", font_size=20, colour=MID_GREY,
            alignment=PP_ALIGN.CENTER)
add_shape_bg(slide, Inches(5.5), Inches(2.7), Inches(2.3), Inches(0.06), ACCENT_BLUE)
add_textbox(slide, Inches(0.8), Inches(6.5), Inches(11.5), Inches(0.5),
            "Confidential — For Informational Purposes Only",
            font_size=12, colour=MID_GREY, alignment=PP_ALIGN.CENTER)


# ── SLIDE 2: Executive Summary ──────────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, WHITE)
section_header(slide, "Executive Summary")

summary_points = [
    "This monthly newsletter provides comprehensive financial analysis of the UK retail sector.",
    "Coverage spans FTSE-listed grocery, fashion, home and specialty retailers.",
    "Key data includes ONS retail sales, BRC-KPMG monitor, consumer confidence, and macroeconomic indicators.",
    "Designed for investors, analysts, and industry professionals tracking UK retail performance.",
]
y = 1.6
for point in summary_points:
    add_textbox(slide, Inches(1.0), Inches(y), Inches(11.3), Inches(0.5),
                f"•  {point}", font_size=16, colour=DARK_GREY)
    y += 0.55

# Key stats boxes
box_data = [
    ("13", "Listed Retailers\nTracked"),
    ("3", "Retail Sectors\nCovered"),
    ("6", "Macro Indicators\nMonitored"),
    ("Monthly", "Publication\nFrequency"),
]
box_x = 1.0
for val, label in box_data:
    add_shape_bg(slide, Inches(box_x), Inches(4.6), Inches(2.6), Inches(2.0), LIGHT_BLUE)
    add_textbox(slide, Inches(box_x), Inches(4.75), Inches(2.6), Inches(0.8),
                val, font_size=36, colour=DARK_BLUE, bold=True, alignment=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(box_x), Inches(5.5), Inches(2.6), Inches(0.7),
                label, font_size=13, colour=DARK_GREY, alignment=PP_ALIGN.CENTER)
    box_x += 2.95


# ── SLIDE 3: Market Overview ────────────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, WHITE)
section_header(slide, "1. Market Overview — FTSE Retail Index Performance")

table_data = [
    ["Metric", "Current", "Previous Month", "YoY Change"],
    ["FTSE 350 General Retailers", "—", "—", "—"],
    ["FTSE 350 Food & Drug Retailers", "—", "—", "—"],
    ["Average Sector P/E Ratio", "—", "—", "—"],
    ["Average Dividend Yield", "—", "—", "—"],
]
add_table(slide, Inches(0.8), Inches(1.6), Inches(11.7), 0.45, table_data,
          col_widths=[4.5, 2.4, 2.4, 2.4])

add_textbox(slide, Inches(0.8), Inches(4.4), Inches(11.5), Inches(0.4),
            "Month-in-Brief", font_size=20, colour=DARK_BLUE, bold=True)
add_textbox(slide, Inches(0.8), Inches(4.9), Inches(11.5), Inches(1.5),
            "•  Summary of key retail events, earnings releases, and market-moving news for March 2026.\n"
            "•  [To be populated with month-specific commentary.]",
            font_size=14, colour=DARK_GREY)


# ── SLIDE 4: KPIs ───────────────────────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, WHITE)
section_header(slide, "2. Key Performance Indicators")

add_textbox(slide, Inches(0.8), Inches(1.4), Inches(5.5), Inches(0.4),
            "UK Retail Sales (ONS Data)", font_size=18, colour=DARK_BLUE, bold=True)
ons_data = [
    ["Indicator", "Latest", "MoM", "YoY"],
    ["Total Retail Sales Volume", "—", "—", "—"],
    ["Food Store Sales", "—", "—", "—"],
    ["Non-Food Store Sales", "—", "—", "—"],
    ["Online Retail (% of total)", "—", "—", "—"],
]
add_table(slide, Inches(0.8), Inches(1.9), Inches(5.5), 0.4, ons_data,
          col_widths=[2.3, 1.1, 1.1, 1.0])

add_textbox(slide, Inches(7.0), Inches(1.4), Inches(5.5), Inches(0.4),
            "BRC-KPMG Retail Sales Monitor", font_size=18, colour=DARK_BLUE, bold=True)
brc_data = [
    ["Metric", "Value", "Trend"],
    ["Total Sales Growth (LFL)", "—", "—"],
    ["Food Sales Growth", "—", "—"],
    ["Non-Food Sales Growth", "—", "—"],
]
add_table(slide, Inches(7.0), Inches(1.9), Inches(5.5), 0.4, brc_data,
          col_widths=[2.5, 1.5, 1.5])


# ── SLIDE 5: Grocery & Food Retail ──────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, WHITE)
section_header(slide, "3a. Listed Retailers — Grocery & Food")

grocery_data = [
    ["Company", "Ticker", "Revenue (£m)", "Rev Growth", "Op. Margin", "Share Price", "MoM Chg"],
    ["Tesco", "TSCO.L", "—", "—", "—", "—", "—"],
    ["Sainsbury's", "SBRY.L", "—", "—", "—", "—", "—"],
    ["Marks & Spencer", "MKS.L", "—", "—", "—", "—", "—"],
    ["Ocado Group", "OCDO.L", "—", "—", "—", "—", "—"],
]
add_table(slide, Inches(0.6), Inches(1.6), Inches(12.1), 0.5, grocery_data,
          col_widths=[2.2, 1.2, 1.8, 1.6, 1.6, 1.8, 1.9])


# ── SLIDE 6: Fashion & General ──────────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, WHITE)
section_header(slide, "3b. Listed Retailers — Fashion & General Merchandise")

fashion_data = [
    ["Company", "Ticker", "Revenue (£m)", "Rev Growth", "Op. Margin", "Share Price", "MoM Chg"],
    ["Next", "NXT.L", "—", "—", "—", "—", "—"],
    ["ABF (Primark)", "ABF.L", "—", "—", "—", "—", "—"],
    ["JD Sports", "JD.L", "—", "—", "—", "—", "—"],
    ["Frasers Group", "FRAS.L", "—", "—", "—", "—", "—"],
    ["boohoo Group", "BOO.L", "—", "—", "—", "—", "—"],
]
add_table(slide, Inches(0.6), Inches(1.6), Inches(12.1), 0.5, fashion_data,
          col_widths=[2.2, 1.2, 1.8, 1.6, 1.6, 1.8, 1.9])


# ── SLIDE 7: Home & Specialty ───────────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, WHITE)
section_header(slide, "3c. Listed Retailers — Home, DIY & Specialty")

home_data = [
    ["Company", "Ticker", "Revenue (£m)", "Rev Growth", "Op. Margin", "Share Price", "MoM Chg"],
    ["Kingfisher (B&Q)", "KGF.L", "—", "—", "—", "—", "—"],
    ["Dunelm", "DNLM.L", "—", "—", "—", "—", "—"],
    ["WHSmith", "SMWH.L", "—", "—", "—", "—", "—"],
    ["Pets at Home", "PETS.L", "—", "—", "—", "—", "—"],
]
add_table(slide, Inches(0.6), Inches(1.6), Inches(12.1), 0.5, home_data,
          col_widths=[2.2, 1.2, 1.8, 1.6, 1.6, 1.8, 1.9])


# ── SLIDE 8: Consumer Spending & Confidence ─────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, WHITE)
section_header(slide, "4. Consumer Spending & Confidence")

add_textbox(slide, Inches(0.8), Inches(1.4), Inches(5.5), Inches(0.4),
            "GfK Consumer Confidence Index", font_size=18, colour=DARK_BLUE, bold=True)
gfk_data = [
    ["Component", "Latest", "Prev Month", "Trend"],
    ["Overall Index", "—", "—", "—"],
    ["Personal Financial Situation", "—", "—", "—"],
    ["Major Purchase Index", "—", "—", "—"],
]
add_table(slide, Inches(0.8), Inches(1.9), Inches(5.5), 0.45, gfk_data,
          col_widths=[2.3, 1.1, 1.1, 1.0])

add_textbox(slide, Inches(7.0), Inches(1.4), Inches(5.5), Inches(0.4),
            "Household Spending Indicators", font_size=18, colour=DARK_BLUE, bold=True)
household_items = [
    ("Avg Weekly Earnings Growth (real)", "—"),
    ("Household Savings Ratio", "—"),
    ("Consumer Credit Growth", "—"),
]
y = 2.0
for label, val in household_items:
    add_shape_bg(slide, Inches(7.0), Inches(y), Inches(5.5), Inches(0.7), LIGHT_BLUE)
    add_textbox(slide, Inches(7.2), Inches(y + 0.1), Inches(3.5), Inches(0.5),
                label, font_size=14, colour=DARK_GREY)
    add_textbox(slide, Inches(10.5), Inches(y + 0.1), Inches(1.8), Inches(0.5),
                val, font_size=16, colour=DARK_BLUE, bold=True, alignment=PP_ALIGN.CENTER)
    y += 0.85


# ── SLIDE 9: Macroeconomic Factors ──────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, WHITE)
section_header(slide, "5. Macroeconomic Factors")

macro_data = [
    ["Indicator", "Latest", "Impact on Retail"],
    ["CPI Inflation (%)", "—", "—"],
    ["BoE Base Rate (%)", "—", "—"],
    ["GBP/USD", "—", "—"],
    ["GBP/EUR", "—", "—"],
    ["UK Unemployment Rate (%)", "—", "—"],
    ["UK GDP Growth (quarterly, %)", "—", "—"],
]
add_table(slide, Inches(0.8), Inches(1.6), Inches(11.7), 0.45, macro_data,
          col_widths=[4.0, 2.0, 5.7])

add_textbox(slide, Inches(0.8), Inches(5.4), Inches(11.5), Inches(0.4),
            "Commentary", font_size=18, colour=DARK_BLUE, bold=True)
add_textbox(slide, Inches(0.8), Inches(5.9), Inches(11.5), Inches(1.2),
            "•  Inflation: [Analysis of how current inflation trends affect retail margins and consumer purchasing power.]\n"
            "•  Interest Rates: [Impact of BoE monetary policy on consumer borrowing and discretionary spending.]\n"
            "•  Currency: [How GBP movements affect import costs for retailers sourcing internationally.]",
            font_size=13, colour=DARK_GREY)


# ── SLIDE 10: E-commerce & Digital Trends ───────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, WHITE)
section_header(slide, "6. E-commerce & Digital Trends")

ecom_data = [
    ["Metric", "Latest", "Trend"],
    ["Online share of total retail (%)", "—", "—"],
    ["IMRG Online Retail Index", "—", "—"],
    ["Mobile commerce share (%)", "—", "—"],
]
add_table(slide, Inches(0.8), Inches(1.6), Inches(6.0), 0.45, ecom_data,
          col_widths=[3.0, 1.5, 1.5])

add_textbox(slide, Inches(0.8), Inches(4.0), Inches(11.5), Inches(0.4),
            "Key Digital Developments", font_size=18, colour=DARK_BLUE, bold=True)
add_textbox(slide, Inches(0.8), Inches(4.5), Inches(11.5), Inches(2.0),
            "•  [Notable digital / e-commerce developments in UK retail this month.]\n"
            "•  [New platform launches, delivery innovations, technology investments.]\n"
            "•  [Changes in online vs in-store shopping behaviour.]",
            font_size=14, colour=DARK_GREY)


# ── SLIDE 11: Sector Risks & Watchlist ──────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, WHITE)
section_header(slide, "7. Sector Risks & Watchlist")

risk_data = [
    ["Risk Factor", "Severity", "Commentary"],
    ["Cost-of-living pressures", "—", "—"],
    ["Supply chain disruption", "—", "—"],
    ["Business rates / regulatory changes", "—", "—"],
    ["Labour market tightness", "—", "—"],
    ["Geopolitical / trade policy risk", "—", "—"],
]
add_table(slide, Inches(0.8), Inches(1.6), Inches(11.7), 0.45, risk_data,
          col_widths=[3.5, 1.7, 6.5])

add_textbox(slide, Inches(0.8), Inches(4.8), Inches(11.5), Inches(0.4),
            "Companies to Watch", font_size=18, colour=DARK_BLUE, bold=True)
add_textbox(slide, Inches(0.8), Inches(5.3), Inches(11.5), Inches(1.5),
            "•  [Company]: [Reason — e.g., upcoming earnings, profit warning, M&A activity.]\n"
            "•  [Company]: [Reason — e.g., management change, strategic review.]\n"
            "•  [Company]: [Reason — e.g., market share shift, new store programme.]",
            font_size=14, colour=DARK_GREY)


# ── SLIDE 12: Outlook & Analyst Consensus ───────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, WHITE)
section_header(slide, "8. Outlook & Analyst Consensus")

outlook_data = [
    ["Metric", "Q2 2026 Est.", "H2 2026 Est.", "Commentary"],
    ["UK Retail Sales Growth (%)", "—", "—", "—"],
    ["Sector EPS Growth (%)", "—", "—", "—"],
    ["Consumer Confidence Direction", "—", "—", "—"],
]
add_table(slide, Inches(0.8), Inches(1.6), Inches(11.7), 0.45, outlook_data,
          col_widths=[3.5, 2.2, 2.2, 3.8])

add_textbox(slide, Inches(0.8), Inches(3.8), Inches(11.5), Inches(0.4),
            "Analyst Commentary", font_size=18, colour=DARK_BLUE, bold=True)
add_shape_bg(slide, Inches(0.8), Inches(4.3), Inches(11.5), Inches(2.2), LIGHT_BLUE)
add_textbox(slide, Inches(1.2), Inches(4.5), Inches(10.7), Inches(1.8),
            "\"[Placeholder for curated analyst quotes and consensus views on UK retail sector outlook "
            "for the coming quarter.]\"\n\n— Analyst Name, Firm",
            font_size=14, colour=DARK_GREY)


# ── SLIDE 13: Data Sources & Disclaimer ─────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, WHITE)
section_header(slide, "Data Sources & Disclaimer")

sources = [
    ("Office for National Statistics (ONS)", "Retail sales, inflation, employment, GDP"),
    ("British Retail Consortium (BRC)", "Retail sales monitor, footfall data"),
    ("GfK", "Consumer confidence barometer"),
    ("IMRG", "Online retail index"),
    ("Bank of England", "Interest rates, credit data"),
    ("London Stock Exchange / LSEG", "Share prices, FTSE indices, analyst estimates"),
]
y = 1.5
for src, desc in sources:
    add_textbox(slide, Inches(1.0), Inches(y), Inches(3.5), Inches(0.35),
                src, font_size=13, colour=DARK_BLUE, bold=True)
    add_textbox(slide, Inches(4.6), Inches(y), Inches(7.0), Inches(0.35),
                desc, font_size=13, colour=DARK_GREY)
    y += 0.4

add_shape_bg(slide, Inches(0.8), Inches(4.5), Inches(11.7), Inches(0.04), MID_GREY)
add_textbox(slide, Inches(0.8), Inches(4.8), Inches(11.5), Inches(1.8),
            "Disclaimer: This newsletter is produced for informational purposes only and does not "
            "constitute financial advice. All data is sourced from publicly available reports and may "
            "be subject to revision. Past performance is not indicative of future results. Readers "
            "should conduct their own research and consult a qualified financial adviser before making "
            "investment decisions.",
            font_size=11, colour=MID_GREY)

add_textbox(slide, Inches(0.8), Inches(6.3), Inches(11.5), Inches(0.5),
            "Next Issue: April 2026  |  For feedback or data requests, contact the editorial team.",
            font_size=13, colour=DARK_BLUE, alignment=PP_ALIGN.CENTER)


# ── Save ─────────────────────────────────────────────────────────────────────
output_path = "newsletters/2026-03_uk-retail-market-review.pptx"
prs.save(output_path)
print(f"Presentation saved: {output_path}")
